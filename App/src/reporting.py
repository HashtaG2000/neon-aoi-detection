from __future__ import annotations

import json
import math
import pathlib
from dataclasses import dataclass
from typing import Iterable

import numpy as np
import pandas as pd


NONE_LABEL = "NoAOI"
RAW_NONE_LABELS = {"", "NoAOI", "nan", "None", "NONE", "null"}
SUMMARY_DIR_NAME = "aoi_master"


@dataclass(frozen=True)
class MasterOutput:
    output_dir: pathlib.Path
    task_metrics_csv: pathlib.Path
    recording_summary_csv: pathlib.Path
    learning_summary_csv: pathlib.Path
    workbook_path: pathlib.Path
    recording_count: int
    task_count: int


def generate_master_outputs(source_dir: pathlib.Path) -> MasterOutput:
    """Build folder-level research summaries from reviewed recording outputs."""
    source_dir = pathlib.Path(source_dir)
    if not source_dir.exists():
        raise FileNotFoundError(f"Source folder does not exist: {source_dir}")

    recording_rows: list[dict[str, object]] = []
    task_rows: list[dict[str, object]] = []

    for rec_dir in _iter_recording_dirs(source_dir):
        final_csv = rec_dir / "aoi_results" / "analysis.csv"
        if not final_csv.exists():
            continue
        df = pd.read_csv(final_csv)
        if df.empty:
            continue

        labels = _label_series(df)
        fps = _estimate_fps(df, len(labels))
        difficulty = _infer_difficulty(rec_dir, source_dir)
        aoi_names = _aoi_names(labels)
        recording_rows.append(_recording_summary_row(rec_dir, difficulty, df, labels, fps, aoi_names))

        tasks = _load_tasks(rec_dir)
        for task_name, bounds in tasks.items():
            row = _task_metrics_row(rec_dir, difficulty, task_name, bounds, df, labels, fps, aoi_names)
            if row:
                task_rows.append(row)

    if not recording_rows:
        raise ValueError(f"No reviewed analysis.csv files found under {source_dir}")

    recording_df = pd.DataFrame(recording_rows).sort_values(["difficulty", "recording"])
    task_df = pd.DataFrame(task_rows)
    if not task_df.empty:
        task_df = task_df.sort_values(["difficulty", "recording", "task_index"])

    learning_df = _learning_summary(task_df)
    aoi_df = _aoi_by_difficulty(task_df)
    quality_df = _data_quality_summary(recording_df, task_df)

    output_dir = source_dir / SUMMARY_DIR_NAME
    output_dir.mkdir(parents=True, exist_ok=True)
    task_csv = output_dir / "master_task_metrics.csv"
    recording_csv = output_dir / "master_recording_summary.csv"
    learning_csv = output_dir / "master_learning_summary.csv"
    workbook_path = output_dir / "master_analysis_workbook.xlsx"

    task_df.to_csv(task_csv, index=False, encoding="utf-8-sig")
    recording_df.to_csv(recording_csv, index=False, encoding="utf-8-sig")
    learning_df.to_csv(learning_csv, index=False, encoding="utf-8-sig")
    _write_workbook(workbook_path, recording_df, task_df, learning_df, aoi_df, quality_df)

    return MasterOutput(
        output_dir=output_dir,
        task_metrics_csv=task_csv,
        recording_summary_csv=recording_csv,
        learning_summary_csv=learning_csv,
        workbook_path=workbook_path,
        recording_count=len(recording_df),
        task_count=len(task_df),
    )


def _iter_recording_dirs(source_dir: pathlib.Path) -> Iterable[pathlib.Path]:
    for path in sorted(source_dir.iterdir()):
        if (
            path.is_dir()
            and not path.name.startswith(".")
            and path.name != SUMMARY_DIR_NAME
            and (path / "info.json").exists()
        ):
            yield path


def _frame_index_array(df: pd.DataFrame) -> np.ndarray | None:
    if "frame_idx" not in df.columns:
        return None
    idx = pd.to_numeric(df["frame_idx"], errors="coerce")
    if idx.isna().any():
        return None
    return idx.astype(int).to_numpy()


def _slice_by_frame_range(
    df: pd.DataFrame,
    labels: pd.Series,
    start_frame: int,
    end_frame: int,
) -> tuple[pd.DataFrame, pd.Series] | None:
    """Slice analysis rows by absolute scene frame_idx (task annotation coordinates)."""
    if end_frame <= start_frame:
        return None
    fi = _frame_index_array(df)
    if fi is not None:
        mask = (fi >= start_frame) & (fi <= end_frame)
        if not mask.any():
            return None
        return df.loc[mask], labels.loc[mask]
    start_frame = max(0, start_frame)
    end_frame = min(end_frame, len(labels) - 1)
    if end_frame <= start_frame:
        return None
    return df.iloc[start_frame : end_frame + 1], labels.iloc[start_frame : end_frame + 1]


def _label_series(df: pd.DataFrame) -> pd.Series:
    source_col = "final_primary_aoi" if "final_primary_aoi" in df.columns else "primary_aoi"
    return df[source_col].map(_normalize_label)


def _normalize_label(value: object) -> str:
    if pd.isna(value):
        return NONE_LABEL
    label = str(value).strip()
    if label in RAW_NONE_LABELS:
        return NONE_LABEL
    return label


def _aoi_names(labels: pd.Series) -> list[str]:
    names = sorted({str(label) for label in labels if str(label) != NONE_LABEL})
    return names


def _estimate_fps(df: pd.DataFrame, frame_count: int) -> float:
    if "time_s" in df.columns and frame_count > 1:
        times = pd.to_numeric(df["time_s"], errors="coerce").dropna()
        if len(times) > 1:
            duration = float(times.iloc[-1] - times.iloc[0])
            if duration > 0:
                return max(1.0, (len(times) - 1) / duration)
    return 30.0


def _infer_difficulty(rec_dir: pathlib.Path, source_dir: pathlib.Path) -> str:
    known = {"easy": "Easy", "medium": "Medium", "hard": "Hard"}
    for part in reversed([*rec_dir.parts, *source_dir.parts]):
        key = part.lower()
        if key in known:
            return known[key]
    return source_dir.name


def _load_tasks(rec_dir: pathlib.Path) -> dict[str, dict[str, int | None]]:
    candidates = [
        rec_dir / "aoi_results" / "raw" / "tasks.json",
        rec_dir / "aoi_results" / "tasks.json",
    ]
    for path in candidates:
        if path.exists():
            try:
                payload = json.loads(path.read_text(encoding="utf-8-sig"))
            except Exception:
                return {}
            if isinstance(payload, dict):
                return {str(k): v for k, v in payload.items() if isinstance(v, dict)}
    return {}


def _recording_summary_row(
    rec_dir: pathlib.Path,
    difficulty: str,
    df: pd.DataFrame,
    labels: pd.Series,
    fps: float,
    aoi_names: list[str],
) -> dict[str, object]:
    frame_count = int(len(labels))
    duration_s = _recording_duration_s(df, frame_count, fps)
    row: dict[str, object] = {
        "recording": rec_dir.name,
        "difficulty": difficulty,
        "frame_count": frame_count,
        "duration_s": duration_s,
        "estimated_fps": fps,
        "none_ratio": _ratio(labels == NONE_LABEL),
        "manual_edit_ratio": _edit_ratio(df, "manual"),
        "auto_gap_fill_ratio": _edit_ratio(df, "auto_gap_fill"),
        "aoi_transition_count": _transition_count(labels),
        "aoi_entropy_bits": _entropy(labels),
    }
    for aoi in aoi_names:
        count = int((labels == aoi).sum())
        row[f"{aoi}_dwell_s"] = count / fps
        row[f"{aoi}_dwell_ratio"] = count / frame_count if frame_count else 0.0
    return row


def _recording_duration_s(df: pd.DataFrame, frame_count: int, fps: float) -> float:
    if "time_s" in df.columns:
        times = pd.to_numeric(df["time_s"], errors="coerce").dropna()
        if len(times) > 1:
            return float(times.iloc[-1] - times.iloc[0])
    return frame_count / fps if fps else 0.0


def _task_metrics_row(
    rec_dir: pathlib.Path,
    difficulty: str,
    task_name: str,
    bounds: dict[str, object],
    df: pd.DataFrame,
    labels: pd.Series,
    fps: float,
    aoi_names: list[str],
) -> dict[str, object] | None:
    start = _safe_int(bounds.get("start"))
    end = _safe_int(bounds.get("end"))
    if start is None or end is None or end <= start:
        return None

    sliced = _slice_by_frame_range(df, labels, start, end)
    if sliced is None:
        return None
    task_df, task_labels = sliced
    frame_count = int(len(task_labels))
    duration_s = _duration_between_frame_range(df, start, end, fps)
    row: dict[str, object] = {
        "recording": rec_dir.name,
        "difficulty": difficulty,
        "task_name": task_name,
        "task_index": _task_index(task_name),
        "start_frame": start,
        "end_frame": end,
        "frame_count": frame_count,
        "duration_s": duration_s,
        "none_ratio": _ratio(task_labels == NONE_LABEL),
        "manual_edit_ratio": _edit_ratio(task_df, "manual"),
        "auto_gap_fill_ratio": _edit_ratio(task_df, "auto_gap_fill"),
        "aoi_transition_count": _transition_count(task_labels),
        "aoi_entropy_bits": _entropy(task_labels),
    }
    for aoi in aoi_names:
        count = int((task_labels == aoi).sum())
        row[f"{aoi}_dwell_s"] = count / fps
        row[f"{aoi}_dwell_ratio"] = count / frame_count if frame_count else 0.0
    row["screen_to_board_ratio"] = _safe_div(row.get("Screen_dwell_s", 0.0), row.get("Board_dwell_s", 0.0))
    row["task_focus_ratio"] = _safe_div(
        float(row.get("Screen_dwell_s", 0.0)) + float(row.get("Board_dwell_s", 0.0)),
        duration_s,
    )
    return row


def _safe_int(value: object) -> int | None:
    try:
        if value is None or pd.isna(value):
            return None
        return int(value)
    except Exception:
        return None


def _task_index(task_name: str) -> int:
    digits = "".join(ch for ch in task_name if ch.isdigit())
    return int(digits) if digits else 0


def _duration_between_frame_range(df: pd.DataFrame, start: int, end: int, fps: float) -> float:
    fi = _frame_index_array(df)
    if fi is not None and "time_s" in df.columns:
        times = pd.to_numeric(df["time_s"], errors="coerce")
        mask = (fi >= start) & (fi <= end)
        subset = times[mask].dropna()
        if len(subset) > 1:
            return max(0.0, float(subset.iloc[-1] - subset.iloc[0]))
    if "time_s" in df.columns and _frame_index_array(df) is None:
        times = pd.to_numeric(df["time_s"], errors="coerce")
        if start < len(times) and end < len(times) and not pd.isna(times.iloc[start]) and not pd.isna(times.iloc[end]):
            return max(0.0, float(times.iloc[end] - times.iloc[start]))
    return max(0.0, (end - start + 1) / fps)


def _ratio(mask: pd.Series | np.ndarray) -> float:
    size = int(len(mask))
    if size == 0:
        return 0.0
    return float(np.asarray(mask, dtype=bool).sum() / size)


def _edit_ratio(df: pd.DataFrame, source: str) -> float:
    if "edit_source" not in df.columns or df.empty:
        return 0.0
    return _ratio(df["edit_source"].astype(str) == source)


def _transition_count(labels: pd.Series) -> int:
    if len(labels) <= 1:
        return 0
    values = labels.astype(str).to_numpy()
    return int(np.sum(values[1:] != values[:-1]))


def _entropy(labels: pd.Series) -> float:
    if labels.empty:
        return 0.0
    counts = labels.value_counts(normalize=True)
    return float(-(counts * np.log2(counts)).sum())


def _safe_div(numerator: object, denominator: object) -> float:
    try:
        den = float(denominator)
        if den == 0:
            return 0.0
        return float(numerator) / den
    except Exception:
        return 0.0


def _learning_summary(task_df: pd.DataFrame) -> pd.DataFrame:
    if task_df.empty:
        return pd.DataFrame()
    grouped = task_df.groupby(["difficulty", "task_index", "task_name"], dropna=False)
    summary = grouped.agg(
        participants=("recording", "nunique"),
        mean_duration_s=("duration_s", "mean"),
        median_duration_s=("duration_s", "median"),
        std_duration_s=("duration_s", "std"),
        mean_none_ratio=("none_ratio", "mean"),
        mean_manual_edit_ratio=("manual_edit_ratio", "mean"),
        mean_transition_count=("aoi_transition_count", "mean"),
        mean_aoi_entropy_bits=("aoi_entropy_bits", "mean"),
        mean_task_focus_ratio=("task_focus_ratio", "mean"),
        mean_screen_to_board_ratio=("screen_to_board_ratio", "mean"),
    ).reset_index()
    summary["sem_duration_s"] = summary["std_duration_s"] / np.sqrt(summary["participants"].clip(lower=1))
    summary["ci95_duration_s"] = 1.96 * summary["sem_duration_s"].fillna(0)
    return summary.sort_values(["difficulty", "task_index"])


def _aoi_by_difficulty(task_df: pd.DataFrame) -> pd.DataFrame:
    if task_df.empty:
        return pd.DataFrame()
    ratio_cols = [col for col in task_df.columns if col.endswith("_dwell_ratio")]
    rows: list[dict[str, object]] = []
    for difficulty, group in task_df.groupby("difficulty"):
        for col in ratio_cols:
            aoi = col[: -len("_dwell_ratio")]
            rows.append(
                {
                    "difficulty": difficulty,
                    "aoi": aoi,
                    "mean_dwell_ratio": float(group[col].fillna(0).mean()),
                    "median_dwell_ratio": float(group[col].fillna(0).median()),
                    "tasks": int(group[col].notna().sum()),
                }
            )
    return pd.DataFrame(rows).sort_values(["difficulty", "aoi"])


def _data_quality_summary(recording_df: pd.DataFrame, task_df: pd.DataFrame) -> pd.DataFrame:
    rec = recording_df.groupby("difficulty", dropna=False).agg(
        recordings=("recording", "nunique"),
        mean_recording_none_ratio=("none_ratio", "mean"),
        mean_recording_manual_edit_ratio=("manual_edit_ratio", "mean"),
        mean_recording_auto_gap_fill_ratio=("auto_gap_fill_ratio", "mean"),
    )
    if task_df.empty:
        return rec.reset_index()
    task = task_df.groupby("difficulty", dropna=False).agg(
        tasks=("task_name", "count"),
        mean_task_none_ratio=("none_ratio", "mean"),
        mean_task_manual_edit_ratio=("manual_edit_ratio", "mean"),
        mean_task_auto_gap_fill_ratio=("auto_gap_fill_ratio", "mean"),
    )
    return rec.join(task, how="outer").reset_index()


def _write_workbook(
    workbook_path: pathlib.Path,
    recording_df: pd.DataFrame,
    task_df: pd.DataFrame,
    learning_df: pd.DataFrame,
    aoi_df: pd.DataFrame,
    quality_df: pd.DataFrame,
) -> None:
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        _readme_df().to_excel(writer, sheet_name="README", index=False)
        recording_df.to_excel(writer, sheet_name="Recording_Summary", index=False)
        task_df.to_excel(writer, sheet_name="Task_Metrics", index=False)
        learning_df.to_excel(writer, sheet_name="Learning_Summary", index=False)
        aoi_df.to_excel(writer, sheet_name="AOI_By_Difficulty", index=False)
        quality_df.to_excel(writer, sheet_name="Data_Quality", index=False)
        _format_workbook(writer.book)


def _readme_df() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"Sheet": "Recording_Summary", "Description": "One row per reviewed recording with dwell, edit, transition, entropy, and data-quality metrics."},
            {"Sheet": "Task_Metrics", "Description": "One row per manually marked task per recording. This is the main statistical-analysis table."},
            {"Sheet": "Learning_Summary", "Description": "Difficulty x task summary with mean/median duration, SEM, and 95% CI."},
            {"Sheet": "AOI_By_Difficulty", "Description": "Mean AOI dwell ratio by difficulty level."},
            {"Sheet": "Data_Quality", "Description": "None/manual-edit/auto-gap-fill rates used to judge data reliability."},
        ]
    )


def _format_workbook(workbook) -> None:
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(color="FFFFFF", bold=True)

    for ws in workbook.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        for col in ws.columns:
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 2, 12), 34)

    learning_ws = workbook["Learning_Summary"]
    if learning_ws.max_row > 1:
        chart = LineChart()
        chart.title = "Mean Task Duration by Task"
        chart.y_axis.title = "Duration (s)"
        chart.x_axis.title = "Task Index"
        data = Reference(learning_ws, min_col=5, min_row=1, max_row=learning_ws.max_row)
        cats = Reference(learning_ws, min_col=2, min_row=2, max_row=learning_ws.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height = 8
        chart.width = 18
        learning_ws.add_chart(chart, "N2")

    quality_ws = workbook["Data_Quality"]
    if quality_ws.max_row > 1 and quality_ws.max_column >= 3:
        chart = BarChart()
        chart.title = "Mean Recording None Ratio"
        chart.y_axis.title = "Ratio"
        chart.x_axis.title = "Difficulty"
        data = Reference(quality_ws, min_col=3, min_row=1, max_row=quality_ws.max_row)
        cats = Reference(quality_ws, min_col=1, min_row=2, max_row=quality_ws.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height = 7
        chart.width = 16
        quality_ws.add_chart(chart, "J2")


# ═══════════════════════════════════════════════════════════════════════════════
#  CONDITION COMPARISON  (NonGamified vs Gamified)
# ═══════════════════════════════════════════════════════════════════════════════

try:
    from scipy import stats as _sp_stats
    _SCIPY_OK = True
except ImportError:
    _SCIPY_OK = False

# Stable AOI display order matching aois.json
_PREFERRED_AOI_ORDER = [
    "Board", "Screen", "Left_Box", "Middle_Box", "Right_Box", "Stream_Deck",
]


@dataclass
class RecordingMetrics:
    """All metrics extracted from one analysed recording."""
    recording:        str
    condition:        str
    dwell_pct:        dict   # aoi -> dwell % of total frames
    fixation_count:   dict   # aoi -> number of fixations
    mean_fix_dur:     dict   # aoi -> mean fixation duration (s); 0 if none
    entropy:          float
    transition_rate:  float  # transitions per second
    trans_matrix:     object  # pd.DataFrame N×N raw counts
    task_durations:   dict   # "Task 1" -> duration_s
    task_aoi_dwell:   dict   # "Task 1" -> {aoi -> dwell_pct}
    task_entropy:     dict   # "Task 1" -> entropy


@dataclass
class ComparisonReport:
    """Pre-computed comparison ready for the dashboard to render."""
    condition_a:    str
    condition_b:    str
    n_a:            int
    n_b:            int
    recordings_a:   list           # list[RecordingMetrics] — for individual data points
    recordings_b:   list           # list[RecordingMetrics] — for individual data points
    aoi_stats:      pd.DataFrame   # per-AOI metrics with stats test results
    overall_stats:  pd.DataFrame   # entropy + transition rate stats
    trans_matrix_a: pd.DataFrame   # row-normalised % averaged across cond A
    trans_matrix_b: pd.DataFrame   # row-normalised % averaged across cond B
    learning_df:    pd.DataFrame   # task_idx / condition / metric / mean / std / n


# ── Helpers ───────────────────────────────────────────────────────────────────

def compute_transition_matrix(
    labels: pd.Series,
    aoi_names: list[str],
) -> pd.DataFrame:
    """Raw count N×N AOI transition matrix (NoAOI excluded as source/dest)."""
    matrix = pd.DataFrame(0, index=aoi_names, columns=aoi_names, dtype=int)
    arr = labels.values
    for i in range(len(arr) - 1):
        src, dst = str(arr[i]), str(arr[i + 1])
        if src != dst and src in matrix.index and dst in matrix.columns:
            matrix.loc[src, dst] += 1
    return matrix


def _row_normalise(matrix: pd.DataFrame) -> pd.DataFrame:
    row_sums = matrix.sum(axis=1)
    return (matrix.div(row_sums.replace(0, np.nan), axis=0).fillna(0) * 100).round(1)


def _mannwhitney(a_vals: list, b_vals: list) -> tuple[float, float]:
    """Returns (p_value, rank-biserial effect size). Both nan if unavailable."""
    if not _SCIPY_OK or len(a_vals) < 2 or len(b_vals) < 2:
        return float("nan"), float("nan")
    try:
        u, p = _sp_stats.mannwhitneyu(a_vals, b_vals, alternative="two-sided")
        effect = 1.0 - (2.0 * float(u)) / (len(a_vals) * len(b_vals))
        return float(p), float(effect)
    except Exception:
        return float("nan"), float("nan")


def _stat_row(
    label: str, a_vals: list, b_vals: list, cond_a: str, cond_b: str
) -> dict:
    p, eff = _mannwhitney(a_vals, b_vals)
    return {
        "metric":              label,
        f"{cond_a}_mean":      round(float(np.mean(a_vals)), 4) if a_vals else float("nan"),
        f"{cond_a}_std":       round(float(np.std(a_vals, ddof=1)), 4) if len(a_vals) > 1 else 0.0,
        f"{cond_b}_mean":      round(float(np.mean(b_vals)), 4) if b_vals else float("nan"),
        f"{cond_b}_std":       round(float(np.std(b_vals, ddof=1)), 4) if len(b_vals) > 1 else 0.0,
        "p_value":             round(p, 4) if not math.isnan(p) else float("nan"),
        "effect_size":         round(eff, 3) if not math.isnan(eff) else float("nan"),
        "significant":         bool(p < 0.05) if not math.isnan(p) else False,
    }


def _load_recording_for_comparison(
    rec_dir: pathlib.Path,
    condition: str,
    aoi_names: list[str],
) -> RecordingMetrics | None:
    """Return RecordingMetrics for one recording, or None if not yet analysed."""
    csv_path = None
    for cand in [
        rec_dir / "aoi_results" / "analysis.csv",
        rec_dir / "aoi_results" / "raw" / "analysis.csv",
    ]:
        if cand.exists() and cand.stat().st_size > 0:
            csv_path = cand
            break
    if csv_path is None:
        return None

    try:
        df = pd.read_csv(csv_path)
    except Exception:
        return None
    if df.empty:
        return None

    labels     = _label_series(df)
    total      = len(labels)
    fps        = _estimate_fps(df, total)
    duration_s = _recording_duration_s(df, total, fps)

    dwell_pct = {
        a: float((labels == a).sum() / total * 100) if total > 0 else 0.0
        for a in aoi_names
    }
    entropy       = _entropy(labels)
    trans_count   = _transition_count(labels)
    trans_rate    = trans_count / duration_s if duration_s > 0 else 0.0
    trans_matrix  = compute_transition_matrix(labels, aoi_names)

    # Fixation metrics (fixation_summary.csv — present only after re-analysis)
    fix_count = {a: 0   for a in aoi_names}
    fix_dur   = {a: 0.0 for a in aoi_names}
    for fix_cand in [
        rec_dir / "aoi_results" / "raw" / "fixation_summary.csv",
        rec_dir / "aoi_results" / "fixation_summary.csv",
    ]:
        if fix_cand.exists() and fix_cand.stat().st_size > 0:
            try:
                fdf = pd.read_csv(fix_cand)
                for a in aoi_names:
                    rows = fdf[fdf["dominant_aoi"] == a]
                    fix_count[a] = int(len(rows))
                    fix_dur[a]   = float(rows["duration_s"].mean()) if len(rows) > 0 else 0.0
            except Exception:
                pass
            break

    # Task data
    tasks          = _load_tasks(rec_dir)
    task_durations = {}
    task_aoi_dwell = {}
    task_entropy   = {}
    for t_name, bounds in tasks.items():
        s = _safe_int(bounds.get("start"))
        e = _safe_int(bounds.get("end"))
        if s is None or e is None or e <= s:
            continue
        sliced = _slice_by_frame_range(df, labels, s, e)
        if sliced is None:
            continue
        _, t_lbl = sliced
        t_total = len(t_lbl)
        if t_total == 0:
            continue
        task_durations[t_name] = _duration_between_frame_range(df, s, e, fps)
        task_aoi_dwell[t_name] = {
            a: float((t_lbl == a).sum() / t_total * 100) for a in aoi_names
        }
        task_entropy[t_name] = _entropy(t_lbl)

    return RecordingMetrics(
        recording=rec_dir.name,
        condition=condition,
        dwell_pct=dwell_pct,
        fixation_count=fix_count,
        mean_fix_dur=fix_dur,
        entropy=entropy,
        transition_rate=trans_rate,
        trans_matrix=trans_matrix,
        task_durations=task_durations,
        task_aoi_dwell=task_aoi_dwell,
        task_entropy=task_entropy,
    )


def _build_aoi_stats(
    recs_a: list, recs_b: list,
    aoi_names: list[str], cond_a: str, cond_b: str,
) -> pd.DataFrame:
    rows = []
    for attr, label in [
        ("dwell_pct",      "Dwell %"),
        ("fixation_count", "Fixation Count"),
        ("mean_fix_dur",   "Mean Fixation Duration (s)"),
    ]:
        for aoi in aoi_names:
            a_vals = [getattr(r, attr).get(aoi, 0.0) for r in recs_a]
            b_vals = [getattr(r, attr).get(aoi, 0.0) for r in recs_b]
            row = _stat_row(label, a_vals, b_vals, cond_a, cond_b)
            row["aoi"] = aoi
            rows.append(row)
    df = pd.DataFrame(rows)
    return df[["aoi", "metric"] + [c for c in df.columns if c not in ("aoi", "metric")]]


def _build_overall_stats(
    recs_a: list, recs_b: list, cond_a: str, cond_b: str,
) -> pd.DataFrame:
    rows = [
        _stat_row("Gaze Entropy (bits)",
                  [r.entropy for r in recs_a],
                  [r.entropy for r in recs_b], cond_a, cond_b),
        _stat_row("Transition Rate (per sec)",
                  [r.transition_rate for r in recs_a],
                  [r.transition_rate for r in recs_b], cond_a, cond_b),
    ]
    return pd.DataFrame(rows)


def _build_avg_trans_matrix(
    recordings: list, aoi_names: list[str],
) -> pd.DataFrame:
    if not recordings:
        return pd.DataFrame(0.0, index=aoi_names, columns=aoi_names)
    normed = [_row_normalise(r.trans_matrix) for r in recordings]
    avg = normed[0].copy().astype(float)
    for m in normed[1:]:
        avg = avg.add(m, fill_value=0.0)
    return (avg / len(normed)).round(1)


def _task_idx(task_name: str) -> int:
    digits = "".join(ch for ch in task_name if ch.isdigit())
    return int(digits) if digits else 0


def _build_learning_df(
    recs_a: list, recs_b: list,
    cond_a: str, cond_b: str,
    aoi_names: list[str],
) -> pd.DataFrame:
    board_aoi  = "Board"  if "Board"  in aoi_names else None
    screen_aoi = "Screen" if "Screen" in aoi_names else None

    rows = []
    for recs, cond in [(recs_a, cond_a), (recs_b, cond_b)]:
        # data[task_idx][metric_key] = [values across recordings]
        data: dict[int, dict[str, list]] = {}
        for rec in recs:
            for t_name, dur in rec.task_durations.items():
                tidx = _task_idx(t_name)
                if tidx == 0:
                    continue
                if tidx not in data:
                    data[tidx] = {
                        "duration_s":      [],
                        "board_dwell_pct": [],
                        "screen_dwell_pct":[],
                        "entropy":         [],
                    }
                data[tidx]["duration_s"].append(dur)
                if board_aoi and t_name in rec.task_aoi_dwell:
                    data[tidx]["board_dwell_pct"].append(
                        rec.task_aoi_dwell[t_name].get(board_aoi, 0.0)
                    )
                if screen_aoi and t_name in rec.task_aoi_dwell:
                    data[tidx]["screen_dwell_pct"].append(
                        rec.task_aoi_dwell[t_name].get(screen_aoi, 0.0)
                    )
                if t_name in rec.task_entropy:
                    data[tidx]["entropy"].append(rec.task_entropy[t_name])

        metric_labels = {
            "duration_s":       "Task Duration (s)",
            "board_dwell_pct":  f"{board_aoi} Dwell %" if board_aoi else "Board Dwell %",
            "screen_dwell_pct": f"{screen_aoi} Dwell %" if screen_aoi else "Screen Dwell %",
            "entropy":          "Gaze Entropy (bits)",
        }
        for tidx, mdata in sorted(data.items()):
            for mkey, mlabel in metric_labels.items():
                vals = mdata.get(mkey, [])
                if not vals:
                    continue
                rows.append({
                    "task_idx":   tidx,
                    "condition":  cond,
                    "metric_key": mkey,
                    "metric":     mlabel,
                    "mean":       float(np.mean(vals)),
                    "std":        float(np.std(vals, ddof=1)) if len(vals) > 1 else 0.0,
                    "n":          len(vals),
                })

    return pd.DataFrame(rows) if rows else pd.DataFrame(
        columns=["task_idx", "condition", "metric_key", "metric", "mean", "std", "n"]
    )


# ── Public API ────────────────────────────────────────────────────────────────

def compare_conditions(
    cond_a_dir: pathlib.Path,
    cond_b_dir: pathlib.Path,
) -> ComparisonReport:
    """Compare two experiment conditions (e.g. NonGamified vs Gamified).

    Directories must contain recording sub-folders with analysed
    aoi_results/analysis.csv files.  Recordings without analysis are silently
    skipped.  Returns a ComparisonReport ready for the dashboard.
    """
    cond_a_dir = pathlib.Path(cond_a_dir)
    cond_b_dir = pathlib.Path(cond_b_dir)
    cond_a = cond_a_dir.name
    cond_b = cond_b_dir.name

    # Derive AOI names from available data, then apply preferred order
    seen: set[str] = set()
    for cdir in [cond_a_dir, cond_b_dir]:
        if not cdir.exists():
            continue
        for rec_dir in _iter_recording_dirs(cdir):
            for cand in [
                rec_dir / "aoi_results" / "analysis.csv",
                rec_dir / "aoi_results" / "raw" / "analysis.csv",
            ]:
                if cand.exists() and cand.stat().st_size > 0:
                    try:
                        # Scan the FULL label column (cheap — one column) so AOIs
                        # that first appear late in a long recording are not missed.
                        header = pd.read_csv(cand, nrows=0).columns
                        col = "final_primary_aoi" if "final_primary_aoi" in header else (
                            "primary_aoi" if "primary_aoi" in header else None)
                        if col is not None:
                            lbl = pd.read_csv(cand, usecols=[col])[col].map(_normalize_label)
                            seen.update(str(v) for v in lbl.unique() if str(v) != NONE_LABEL)
                    except Exception:
                        pass
                    break

    aoi_names = [a for a in _PREFERRED_AOI_ORDER if a in seen] + \
                sorted(a for a in seen if a not in _PREFERRED_AOI_ORDER)

    recs_a: list[RecordingMetrics] = []
    if cond_a_dir.exists():
        for rec_dir in _iter_recording_dirs(cond_a_dir):
            m = _load_recording_for_comparison(rec_dir, cond_a, aoi_names)
            if m is not None:
                recs_a.append(m)

    recs_b: list[RecordingMetrics] = []
    if cond_b_dir.exists():
        for rec_dir in _iter_recording_dirs(cond_b_dir):
            m = _load_recording_for_comparison(rec_dir, cond_b, aoi_names)
            if m is not None:
                recs_b.append(m)

    return ComparisonReport(
        condition_a=cond_a,
        condition_b=cond_b,
        n_a=len(recs_a),
        n_b=len(recs_b),
        recordings_a=recs_a,
        recordings_b=recs_b,
        aoi_stats=_build_aoi_stats(recs_a, recs_b, aoi_names, cond_a, cond_b),
        overall_stats=_build_overall_stats(recs_a, recs_b, cond_a, cond_b),
        trans_matrix_a=_build_avg_trans_matrix(recs_a, aoi_names),
        trans_matrix_b=_build_avg_trans_matrix(recs_b, aoi_names),
        learning_df=_build_learning_df(recs_a, recs_b, cond_a, cond_b, aoi_names),
    )
